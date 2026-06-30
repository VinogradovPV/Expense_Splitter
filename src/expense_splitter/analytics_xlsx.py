from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from expense_splitter.analytics import AnalyticsDataset
from expense_splitter.report_tables import chart_display_title, rows_for_visual_table
from expense_splitter.ui_labels import label_for_report_sheet

SHEETS = (
    (label_for_report_sheet("Summary"), "summary.csv"),
    (label_for_report_sheet("Purchases"), "purchases.csv"),
    (label_for_report_sheet("By Category"), "by_category.csv"),
    (label_for_report_sheet("By Payer"), "by_payer.csv"),
    (label_for_report_sheet("By Participant"), "by_participant.csv"),
    (label_for_report_sheet("Balances"), "balances.csv"),
    (label_for_report_sheet("Settlements"), "settlements.csv"),
    (label_for_report_sheet("Top Purchases"), "top_purchases.csv"),
    (label_for_report_sheet("Warnings"), "warnings.csv"),
)

MONEY_HEADERS = {"Сумма", "Оплачено", "Доля расходов", "Доля", "Баланс", "Доля оплат, %"}
INTEGER_HEADERS = {"Количество покупок", "Покупок", "Место"}
SUMMARY_MONEY_LABELS = {"Общая сумма", "Средняя покупка"}
SUMMARY_INTEGER_LABELS = {"Количество покупок", "Количество участников"}

PRIMARY = "285F8F"
SECONDARY = "DCE8F2"
HEADER = "EAF1F7"
INK = "18212F"
MUTED = "64748B"
WARNING = "FFF4D6"
THIN = Side(style="thin", color="DCE3EC")


def write_xlsx_report(
    dataset: AnalyticsDataset,
    path: Path,
    tables_dir: Path,
    charts_dir: Path,
) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)

    for sheet_name, filename in SHEETS:
        rows = rows_for_visual_table(filename, _read_csv(tables_dir / filename))
        worksheet = workbook.create_sheet(sheet_name)
        _write_table_sheet(worksheet, sheet_name, rows)

    charts_sheet = workbook.create_sheet(label_for_report_sheet("Charts"))
    _write_charts_sheet(charts_sheet, charts_dir)
    workbook.active = 0
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _read_csv(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.reader(stream))


def _write_table_sheet(worksheet, sheet_name: str, rows: list[list[str]]) -> None:
    worksheet.sheet_view.showGridLines = False
    max_columns = max((len(row) for row in rows), default=2)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_columns)
    title = worksheet.cell(1, 1, _sheet_title(sheet_name))
    title.font = Font(size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor=PRIMARY)
    title.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 28

    if not rows:
        worksheet.cell(3, 1, "Нет данных.")
        return

    headers = rows[0]
    for column, header in enumerate(headers, 1):
        cell = worksheet.cell(3, column, header)
        cell.font = Font(bold=True, color=INK)
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.border = Border(bottom=THIN)
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row_index, row in enumerate(rows[1:], 4):
        for column, raw_value in enumerate(row, 1):
            header = headers[column - 1]
            value = _typed_value(raw_value, header, row)
            cell = worksheet.cell(row_index, column, value)
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(
                vertical="top", wrap_text=header in {"Комментарий", "Сообщение"}
            )
            if _is_money(header, row):
                cell.number_format = "#,##0.00"
            elif header == "Дата" and isinstance(value, date):
                cell.number_format = "dd.mm.yyyy"

    last_row = max(3, len(rows) + 2)
    last_column = get_column_letter(max_columns)
    worksheet.freeze_panes = "A4"
    worksheet.auto_filter.ref = f"A3:{last_column}{last_row}"
    for row_index in range(4, last_row + 1):
        if row_index % 2 == 0:
            for cell in worksheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")

    if sheet_name == label_for_report_sheet("Warnings"):
        for row in worksheet.iter_rows(min_row=4, max_row=last_row):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=WARNING)
                cell.font = Font(color="7A5313")

    _set_widths(worksheet, headers, rows[1:])


def _typed_value(raw_value: str, header: str, row: list[str]):
    if raw_value == "":
        return None
    if header == "payer_rank":
        try:
            return int(raw_value)
        except ValueError:
            return raw_value
    if header == "Дата":
        try:
            return date.fromisoformat(raw_value)
        except ValueError:
            return raw_value
    if _is_money(header, row):
        try:
            return float(raw_value)
        except ValueError:
            return raw_value
    if header in INTEGER_HEADERS or (row and row[0] in SUMMARY_INTEGER_LABELS):
        try:
            return int(raw_value)
        except ValueError:
            return raw_value
    return raw_value


def _is_money(header: str, row: list[str]) -> bool:
    if header == "payer_total":
        return True
    return header in MONEY_HEADERS or (
        header == "Значение" and row and row[0] in SUMMARY_MONEY_LABELS
    )


def _set_widths(worksheet, headers: list[str], rows: list[list[str]]) -> None:
    for index, header in enumerate(headers, 1):
        values = [header] + [row[index - 1] for row in rows if index <= len(row)]
        width = min(max(max((len(str(value)) for value in values), default=8) + 2, 12), 38)
        worksheet.column_dimensions[get_column_letter(index)].width = width


def _write_charts_sheet(worksheet, charts_dir: Path) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:J1")
    worksheet["A1"] = "Графики аналитики"
    worksheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor=PRIMARY)
    worksheet.row_dimensions[1].height = 28
    chart_paths = sorted(charts_dir.glob("*.png")) if charts_dir.exists() else []
    if not chart_paths:
        worksheet["A3"] = "Графики не созданы: нет данных за выбранный период."
        worksheet["A3"].font = Font(italic=True, color=MUTED)
        worksheet["A3"].fill = PatternFill("solid", fgColor=SECONDARY)
        return

    row = 3
    for chart_path in chart_paths:
        worksheet.cell(row, 1, chart_display_title(chart_path))
        worksheet.cell(row, 1).font = Font(bold=True, color=PRIMARY)
        row += 1
        image = Image(chart_path)
        if image.width > 900:
            ratio = 900 / image.width
            image.width = 900
            image.height = int(image.height * ratio)
        worksheet.add_image(image, f"A{row}")
        row += max(24, int(image.height / 20) + 3)
    for column in range(1, 11):
        worksheet.column_dimensions[get_column_letter(column)].width = 12


def _sheet_title(sheet_name: str) -> str:
    titles = {
        label_for_report_sheet("Summary"): "Сводка аналитики",
        label_for_report_sheet("Purchases"): "Покупки",
        label_for_report_sheet("By Category"): "Расходы по категориям",
        label_for_report_sheet("By Payer"): "Расходы по плательщикам",
        label_for_report_sheet("By Participant"): "Объем расходов на человека",
        label_for_report_sheet("Balances"): "Балансы участников",
        label_for_report_sheet("Settlements"): "Переводы",
        label_for_report_sheet("Top Purchases"): "Крупнейшие покупки",
        label_for_report_sheet("Warnings"): "Предупреждения",
    }
    return titles[sheet_name]
