from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from html import escape
from pathlib import Path
from typing import Iterable, Sequence

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from expense_splitter.analytics import (
    UNCATEGORIZED,
    aggregate_by_category,
    aggregate_by_participant,
    aggregate_by_payer,
    build_top_purchases,
)
from expense_splitter.calculator import calculate_balances
from expense_splitter.models import (
    DEFAULT_PURCHASE_NAME,
    Balance,
    Participant,
    Purchase,
    Settlement,
)
from expense_splitter.report_pdf import PdfTableSpec, write_pdf_report
from expense_splitter.report_sorting import sorted_purchases_with_payer_totals
from expense_splitter.report_tables import (
    BALANCE_CHART_TITLE,
    BALANCE_CHART_XLABEL,
    BALANCE_EXPLANATION,
    BALANCE_RESULT_HEADER,
    PARTICIPANT_LABEL_MAX_LENGTH,
    PURCHASE_LABEL_MAX_LENGTH,
    chart_display_title,
    rows_for_visual_table,
    truncate_display_label,
    visual_table_note,
)
from expense_splitter.settlement import calculate_settlements
from expense_splitter.settlement_periods import filter_purchases_by_settlement_scope
from expense_splitter.ui_labels import (
    label_for_purchase_status,
    label_for_report_field,
    label_for_report_sheet,
    label_for_scope,
)
from expense_splitter.visual.palette import (
    EXPENSE_DIMENSION_COLOR_MAP,
    PALETTE_NAME,
    PALETTE_VERSION,
    QUALITATIVE_PALETTE,
    build_period_color_map,
    build_stable_color_map,
    color_for_balance_status,
)

SUPPORTED_FORMATS = {"markdown", "csv", "png", "html", "xlsx", "pdf", "all"}
SUPPORTED_SCOPES = {"open", "all"}
CSV_ENCODING = "utf-8-sig"
EMPTY_OPEN_SCOPE_MESSAGE = "Нет открытых покупок для текущих взаиморасчетов."
OPERATIONS_BY_DAY_FILENAME = "operations_by_day.png"
OPERATIONS_BY_DAY_WARNING_TYPE = "chart_not_enough_data"
OPERATIONS_BY_DAY_WARNING_MESSAGE = "Недостаточно дат для построения графика операций по дням."

TABLE_SPECS = (
    ("summary.csv", "Сводка"),
    ("settlements.csv", "Итоговые переводы"),
    ("purchases.csv", "Покупки"),
    ("by_category.csv", "Расходы по категориям"),
    ("by_payer.csv", "Расходы по плательщикам"),
    ("by_participant.csv", "Объем расходов на человека"),
    ("balances.csv", "Балансы"),
    ("warnings.csv", "Предупреждения"),
)

CHART_SPECS = (
    ("balances.png", chart_display_title("balances.png")),
    ("participant_share.png", chart_display_title("participant_share.png")),
    (OPERATIONS_BY_DAY_FILENAME, chart_display_title(OPERATIONS_BY_DAY_FILENAME)),
    ("spending_by_category.png", "Расходы по категориям"),
    ("spending_by_payer.png", "Расходы по плательщикам"),
    ("top_purchases.png", "Крупнейшие покупки"),
)

CURRENT_PDF_TABLES = (
    PdfTableSpec(
        "settlements.csv",
        "Итоговые переводы",
        display_mode="financial",
        section="overview",
    ),
    PdfTableSpec("by_payer.csv", "Расходы по плательщикам", section="overview"),
    PdfTableSpec("by_category.csv", "Расходы по категориям", section="overview"),
    PdfTableSpec("warnings.csv", "Предупреждения", display_mode="callout", section="overview"),
    PdfTableSpec("by_participant.csv", "Объем расходов на человека", section="participants"),
    PdfTableSpec("balances.csv", "Балансы", display_mode="financial", section="participants"),
    PdfTableSpec("purchases.csv", "Покупки", display_mode="full", section="participants"),
)

CURRENT_HTML_CSS = """
:root {
  --ink:#172033; --muted:#607086; --line:#d9e2ec; --paper:#f6f8fb;
  --card:#fff; --accent:#285f8f; --warn:#8a5a13;
}
* { box-sizing:border-box; }
body {
  margin:0; background:var(--paper); color:var(--ink);
  font:15px/1.55 system-ui,-apple-system,"Segoe UI",sans-serif;
}
main { max-width:1180px; margin:auto; padding:30px 20px 64px; }
h1,h2,h3 { line-height:1.2; }
h2 { margin-top:34px; }
.meta,.muted { color:var(--muted); }
.cards {
  display:grid; grid-template-columns:repeat(auto-fit,minmax(170px,1fr));
  gap:12px; margin:22px 0;
}
.card,.panel {
  background:var(--card); border:1px solid var(--line); border-radius:8px;
}
.card { padding:16px; }
.card span { color:var(--muted); }
.card strong { display:block; font-size:24px; margin-top:4px; }
.panel { padding:18px; margin:14px 0; overflow:auto; }
.warning {
  background:#fff9ed; border-left:4px solid #d99a32; color:var(--warn);
}
.empty { padding:14px; background:#eef5fb; border:1px solid var(--line); }
table { width:100%; min-width:640px; border-collapse:collapse; }
th,td { padding:8px 10px; border-bottom:1px solid var(--line); text-align:left; }
td.num { text-align:right; font-variant-numeric:tabular-nums; }
th { background:#eef3f8; }
.chart { display:block; max-width:100%; height:auto; }
.missing { color:var(--muted); border:1px dashed var(--line); padding:14px; }
.links { display:flex; flex-wrap:wrap; gap:12px; }
a { color:var(--accent); }
"""


@dataclass(frozen=True)
class CurrentReportDataset:
    scope: str
    generated_at: datetime
    participants: list[str]
    purchases: list[Purchase]
    balances: list[Balance]
    settlements: list[Settlement]
    summary: dict[str, object]
    by_category: list[dict[str, object]]
    by_payer: list[dict[str, object]]
    by_participant: list[dict[str, object]]
    top_purchases: list[dict[str, object]]
    warnings: list[dict[str, object]]


@dataclass(frozen=True)
class DailyOperationPoint:
    date: date
    total_amount: Decimal
    purchase_count: int


def build_current_report_dataset(
    participants: Sequence[Participant | str],
    purchases: Sequence[Purchase],
    scope: str = "open",
    generated_at: datetime | None = None,
) -> CurrentReportDataset:
    scope_key = scope.lower()
    if scope_key not in SUPPORTED_SCOPES:
        raise ValueError("Scope must be one of: open, all.")

    participant_names = _participant_names(participants)
    selected = filter_purchases_by_settlement_scope(purchases, scope=scope_key)  # type: ignore[arg-type]
    balances = calculate_balances(selected, participant_names)
    settlements = calculate_settlements(balances)
    generated = generated_at or datetime.now(timezone.utc).astimezone()

    summary = _summary(selected, participant_names, settlements, scope_key, generated)
    by_category = _with_share_percent(aggregate_by_category(selected), summary["total_amount"])

    return CurrentReportDataset(
        scope=scope_key,
        generated_at=generated,
        participants=participant_names,
        purchases=selected,
        balances=balances,
        settlements=settlements,
        summary=summary,
        by_category=by_category,
        by_payer=aggregate_by_payer(selected),
        by_participant=aggregate_by_participant(selected, participant_names),
        top_purchases=build_top_purchases(selected),
        warnings=_warnings(selected, participant_names, scope_key),
    )


def generate_current_report(
    dataset: CurrentReportDataset,
    output_root: Path = Path("reports/current_state"),
    output_format: str = "all",
) -> Path:
    format_key = output_format.lower()
    if format_key not in SUPPORTED_FORMATS:
        raise ValueError("Format must be one of: markdown, csv, png, html, xlsx, pdf, all.")

    report_dir = output_root / f"{dataset.scope}_{dataset.generated_at:%Y-%m-%d_%H-%M-%S}"
    tables_dir = report_dir / "tables"
    charts_dir = report_dir / "charts"
    tables_dir.mkdir(parents=True, exist_ok=True)
    charts_dir.mkdir(parents=True, exist_ok=True)

    warnings = [dict(row) for row in dataset.warnings]
    generated: list[Path] = []
    chart_paths: list[Path] = []

    if format_key in {"csv", "html", "xlsx", "pdf", "all"}:
        generated.extend(write_current_csv_tables(dataset, tables_dir, warnings))
    else:
        generated.extend(write_minimum_tables(dataset, tables_dir, warnings))

    if format_key in {"png", "html", "xlsx", "pdf", "all"}:
        chart_paths, chart_warnings = write_current_charts(dataset, charts_dir)
        generated.extend(chart_paths)
        warnings.extend(chart_warnings)
        write_warnings_table(tables_dir / "warnings.csv", warnings)

    if format_key in {"markdown", "html", "all"}:
        generated.append(
            write_current_markdown(
                dataset,
                report_dir / "current_state_report.md",
                warnings,
                chart_paths=chart_paths,
            )
        )

    if format_key in {"html", "all"}:
        generated.append(
            write_current_html(
                dataset,
                report_dir / "current_state_dashboard.html",
                warnings,
                chart_paths=chart_paths,
            )
        )

    if format_key in {"xlsx", "all"}:
        generated.append(
            write_current_xlsx(
                dataset,
                report_dir / "current_state.xlsx",
                tables_dir,
                charts_dir,
                chart_paths=chart_paths,
            )
        )

    if format_key in {"pdf", "all"}:
        generated.append(
            write_pdf_report(
                report_dir / "current_state.pdf",
                title="Отчет по текущим взаиморасчетам",
                metadata_rows=_summary_rows(dataset),
                tables_dir=tables_dir,
                table_specs=(
                    *CURRENT_PDF_TABLES,
                    *(
                        (
                            PdfTableSpec(
                                "purchases.csv",
                                "Все покупки",
                                display_mode="appendix",
                                section="appendix",
                            ),
                        )
                        if len(dataset.purchases) > 15
                        else ()
                    ),
                ),
                charts_dir=charts_dir,
                chart_paths=chart_paths,
                report_kind="current",
            )
        )

    metadata_path = report_dir / "metadata.json"
    generated.append(metadata_path)
    write_current_metadata(dataset, metadata_path, generated, warnings, report_dir, chart_paths)
    return report_dir


def write_current_csv_tables(
    dataset: CurrentReportDataset,
    tables_dir: Path,
    warnings: Sequence[dict[str, object]] | None = None,
) -> list[Path]:
    warning_rows = list(dataset.warnings if warnings is None else warnings)
    specs = (
        ("summary.csv", ["Показатель", "Значение"], _summary_rows(dataset)),
        (
            "purchases.csv",
            [
                "ID",
                "Дата",
                "Покупка",
                "Категория",
                "Сумма",
                "Плательщик",
                "Участники",
                "Статус",
                "Комментарий",
            ],
            _purchase_rows(dataset),
        ),
        (
            "by_category.csv",
            ["Категория", "Сумма", "Покупок", "Доля, %"],
            (
                [
                    row["category"],
                    row["total_amount"],
                    row["purchase_count"],
                    row["share_percent"],
                ]
                for row in dataset.by_category
            ),
        ),
        (
            "by_payer.csv",
            ["Плательщик", "Оплачено", "Покупок", "Доля оплат, %"],
            (
                [
                    row["payer"],
                    row["total_paid"],
                    row["purchase_count"],
                    row["payer_share_percent"],
                ]
                for row in dataset.by_payer
            ),
        ),
        (
            "by_participant.csv",
            ["Участник", "Доля расходов", "Покупок"],
            (
                [row["participant"], row["total_share"], row["purchase_count"]]
                for row in dataset.by_participant
            ),
        ),
        (
            "balances.csv",
            ["Участник", "Оплачено", "Доля", "Баланс"],
            ([row.participant, row.paid, row.share, row.net] for row in dataset.balances),
        ),
        (
            "settlements.csv",
            ["От кого", "Кому", "Сумма"],
            ([row.from_participant, row.to_participant, row.amount] for row in dataset.settlements),
        ),
        (
            "warnings.csv",
            ["Тип", "ID покупки", "Покупка", "Сообщение"],
            _warning_rows(warning_rows),
        ),
    )
    paths: list[Path] = []
    for filename, headers, rows in specs:
        path = tables_dir / filename
        _write_csv(path, headers, rows)
        paths.append(path)
    return paths


def write_minimum_tables(
    dataset: CurrentReportDataset,
    tables_dir: Path,
    warnings: Sequence[dict[str, object]],
) -> list[Path]:
    paths: list[Path] = []
    summary_path = tables_dir / "summary.csv"
    _write_csv(summary_path, ["Показатель", "Значение"], _summary_rows(dataset))
    paths.append(summary_path)
    warning_path = tables_dir / "warnings.csv"
    write_warnings_table(warning_path, warnings)
    paths.append(warning_path)
    return paths


def write_warnings_table(path: Path, warnings: Sequence[dict[str, object]]) -> Path:
    _write_csv(path, ["Тип", "ID покупки", "Покупка", "Сообщение"], _warning_rows(warnings))
    return path


def write_current_markdown(
    dataset: CurrentReportDataset,
    path: Path,
    warnings: Sequence[dict[str, object]],
    chart_paths: Sequence[Path] = (),
) -> Path:
    summary = dataset.summary
    lines = [
        "# Отчет по текущим взаиморасчетам",
        "",
        f"{label_for_report_field('Scope')}: {label_for_scope(dataset.scope)}",
        (
            f"{label_for_report_field('Snapshot')}: "
            f"{dataset.generated_at.isoformat(timespec='seconds')}"
        ),
        "",
        "## Сводка",
        "",
        "| Показатель | Значение |",
        "|---|---:|",
        f"| Общая сумма | {_serialize(summary['total_amount'])} |",
        f"| Количество покупок | {summary['purchase_count']} |",
        f"| Средний чек | {_serialize(summary['average_purchase'])} |",
        f"| Количество участников | {summary['participant_count']} |",
        f"| Итоговые переводы | {summary['settlement_count']} |",
        "",
    ]
    if not dataset.purchases:
        message = EMPTY_OPEN_SCOPE_MESSAGE if dataset.scope == "open" else "Нет покупок для отчета."
        lines.extend([message, ""])

    _append_markdown_table(
        lines,
        "Итоговые переводы",
        ["От кого", "Кому", "Сумма"],
        ([row.from_participant, row.to_participant, row.amount] for row in dataset.settlements),
    )
    _append_markdown_table(
        lines,
        "Расходы по категориям",
        ["Категория", "Сумма", "Покупок", "Доля, %"],
        (
            [row["category"], row["total_amount"], row["purchase_count"], row["share_percent"]]
            for row in dataset.by_category
        ),
    )
    _append_markdown_table(
        lines,
        "Расходы по плательщикам",
        ["Плательщик", "Оплачено", "Покупок", "Доля оплат, %"],
        (
            [row["payer"], row["total_paid"], row["purchase_count"], row["payer_share_percent"]]
            for row in dataset.by_payer
        ),
    )
    _append_markdown_table(
        lines,
        "Объем расходов на человека",
        ["Участник", "Объем расходов на человека", "Покупок"],
        (
            [row["participant"], row["total_share"], row["purchase_count"]]
            for row in dataset.by_participant
        ),
    )
    _append_markdown_table(
        lines,
        "Балансы",
        ["Участник", "Оплачено", "Объем расходов на человека", BALANCE_RESULT_HEADER],
        ([row.participant, row.paid, row.share, row.net] for row in dataset.balances),
    )
    lines.extend([BALANCE_EXPLANATION, ""])
    _append_markdown_table(
        lines,
        "Покупки в текущем расчете",
        [
            "Дата",
            "Покупка",
            "Категория",
            "Сумма",
            "Плательщик",
            "Участники",
            "Статус",
            "Комментарий",
        ],
        (
            [row[1], row[2], row[3], row[4], row[5], row[6], row[7], row[8]]
            for row in _purchase_rows(dataset)
        ),
    )
    _append_markdown_table(
        lines,
        "Предупреждения",
        ["Тип", "ID покупки", "Покупка", "Сообщение"],
        _warning_rows(warnings),
    )

    current_chart_paths = sorted(chart_paths, key=_chart_sort_key)
    if current_chart_paths:
        lines.extend(["## Графики", ""])
        for chart_path in current_chart_paths:
            title = chart_display_title(chart_path)
            lines.extend([f"![{title}](charts/{chart_path.name})", ""])

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_current_html(
    dataset: CurrentReportDataset,
    path: Path,
    warnings: Sequence[dict[str, object]],
    chart_paths: Sequence[Path] = (),
) -> Path:
    summary = dataset.summary
    cards = (
        ("Общая сумма", _serialize(summary["total_amount"])),
        ("Покупок", str(summary["purchase_count"])),
        ("Средний чек", _serialize(summary["average_purchase"])),
        ("Участников", str(summary["participant_count"])),
        ("Переводов", str(summary["settlement_count"])),
    )
    card_html = "".join(
        f'<div class="card"><span>{escape(label)}</span><strong>{escape(value)}</strong></div>'
        for label, value in cards
    )
    table_html = "".join(
        _html_table(path.parent, filename, title) for filename, title in TABLE_SPECS
    )
    current_chart_paths = sorted(chart_paths, key=_chart_sort_key)
    chart_html = "".join(_html_chart(chart_path) for chart_path in current_chart_paths)
    warning_items = "".join(f"<li>{escape(str(row.get('message', '')))}</li>" for row in warnings)
    warnings_html = f"<ul>{warning_items}</ul>" if warning_items else "<p>Предупреждений нет.</p>"
    csv_links = "".join(
        f'<a href="tables/{escape(filename)}">{escape(title)} CSV</a>'
        for filename, title in TABLE_SPECS
    )
    empty_note = (
        f'<p class="empty">{escape(EMPTY_OPEN_SCOPE_MESSAGE)}</p>'
        if not dataset.purchases and dataset.scope == "open"
        else ""
    )
    scope_label = escape(label_for_report_field("Scope"))
    scope_value = escape(label_for_scope(dataset.scope))
    snapshot_label = escape(label_for_report_field("Snapshot"))
    snapshot_value = escape(dataset.generated_at.isoformat(timespec="seconds"))
    document = f"""<!doctype html>
<html lang="ru"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Отчет по текущим взаиморасчетам</title>
<style>{CURRENT_HTML_CSS}</style></head><body><main>
<header><h1>Отчет по текущим взаиморасчетам</h1>
<p class="meta">{scope_label}: {scope_value}<br>
{snapshot_label}: {snapshot_value}</p></header>
<section><h2>Сводка</h2>{empty_note}<div class="cards">{card_html}</div></section>
<section><h2>Предупреждения</h2><div class="panel warning">{warnings_html}</div></section>
<section><h2>Графики</h2>{chart_html}</section>
<section><h2>Таблицы</h2>{table_html}</section>
<section><h2>Файлы</h2><div class="panel links">
<a href="current_state_report.md">Markdown</a>{csv_links}</div></section>
</main></body></html>"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(document, encoding="utf-8")
    return path


def write_current_xlsx(
    dataset: CurrentReportDataset,
    path: Path,
    tables_dir: Path,
    charts_dir: Path,
    chart_paths: Sequence[Path] = (),
) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = (
        (label_for_report_sheet("Summary"), "summary.csv"),
        (label_for_report_sheet("Settlements"), "settlements.csv"),
        (label_for_report_sheet("Purchases"), "purchases.csv"),
        (label_for_report_sheet("By Category"), "by_category.csv"),
        (label_for_report_sheet("By Payer"), "by_payer.csv"),
        (label_for_report_sheet("By Participant"), "by_participant.csv"),
        (label_for_report_sheet("Balances"), "balances.csv"),
        (label_for_report_sheet("Warnings"), "warnings.csv"),
    )
    for sheet_name, filename in sheets:
        worksheet = workbook.create_sheet(sheet_name)
        _write_sheet(
            worksheet,
            sheet_name,
            rows_for_visual_table(filename, _read_csv(tables_dir / filename)),
        )
    charts_sheet = workbook.create_sheet(label_for_report_sheet("Charts"))
    _write_charts_sheet(charts_sheet, charts_dir, chart_paths)
    workbook.active = 0
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def write_current_charts(
    dataset: CurrentReportDataset,
    output_dir: Path,
) -> tuple[list[Path], list[dict[str, object]]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    _clean_known_chart_outputs(output_dir)
    generated: list[Path] = []
    warnings: list[dict[str, object]] = []
    specs = (
        ("balances.png", dataset.balances, _chart_balances),
        ("participant_share.png", dataset.by_participant, _chart_participant_share),
        (
            OPERATIONS_BY_DAY_FILENAME,
            aggregate_daily_operations(dataset.purchases),
            _chart_operations_by_day,
        ),
        ("spending_by_category.png", dataset.by_category, _chart_category),
        ("spending_by_payer.png", dataset.by_payer, _chart_payer),
        ("top_purchases.png", dataset.top_purchases, _chart_top_purchases),
    )
    for filename, rows, renderer in specs:
        if filename == OPERATIONS_BY_DAY_FILENAME and dataset.purchases and len(rows) < 2:
            warnings.append(_operations_by_day_not_enough_data_warning())
            continue
        if not dataset.purchases or not rows:
            warnings.append(_chart_warning(filename))
            continue
        path = output_dir / filename
        renderer(dataset, path)
        generated.append(path)
    return generated, warnings


def _clean_known_chart_outputs(output_dir: Path) -> None:
    for filename, _title in CHART_SPECS:
        path = output_dir / filename
        if path.exists():
            path.unlink()


def write_current_metadata(
    dataset: CurrentReportDataset,
    path: Path,
    generated_files: Sequence[Path],
    warnings: Sequence[dict[str, object]],
    report_dir: Path,
    chart_paths: Sequence[Path] = (),
) -> Path:
    generated_chart_keys = [
        chart_path.stem for chart_path in sorted(chart_paths, key=_chart_sort_key)
    ]
    metadata = {
        "schema_version": 1,
        "report_type": "current_state",
        "scope": dataset.scope,
        "generated_at": dataset.generated_at.isoformat(),
        "summary": {key: _json_value(value) for key, value in dataset.summary.items()},
        "palette": {"name": PALETTE_NAME, "version": PALETTE_VERSION},
        "warning_count": len(warnings),
        "warnings": [dict(row) for row in warnings],
        "charts_generated": generated_chart_keys,
        "charts_skipped": [
            {
                "chart": row.get("chart"),
                "reason": row.get("warning_type"),
                "message": row.get("message"),
            }
            for row in warnings
            if row.get("chart")
            and row.get("warning_type") in {"chart_no_data", "chart_not_enough_data"}
        ],
        "files": sorted(
            str(item.relative_to(report_dir)).replace("\\", "/")
            for item in generated_files
            if item.exists()
        ),
    }
    path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def _summary(
    purchases: Sequence[Purchase],
    participants: Sequence[str],
    settlements: Sequence[Settlement],
    scope: str,
    generated_at: datetime,
) -> dict[str, object]:
    total_amount = sum((purchase.amount for purchase in purchases), Decimal("0.00"))
    purchase_count = len(purchases)
    average_purchase = (
        (total_amount / Decimal(purchase_count)).quantize(Decimal("0.01"))
        if purchase_count
        else Decimal("0.00")
    )
    return {
        "scope": scope,
        "generated_at": generated_at.isoformat(timespec="seconds"),
        "total_amount": total_amount,
        "purchase_count": purchase_count,
        "average_purchase": average_purchase,
        "participant_count": len(participants),
        "settlement_count": len(settlements),
    }


def _with_share_percent(
    rows: list[dict[str, object]],
    total_amount: object,
) -> list[dict[str, object]]:
    total = total_amount if isinstance(total_amount, Decimal) else Decimal(str(total_amount))
    result: list[dict[str, object]] = []
    for row in rows:
        amount = row["total_amount"]
        value = amount if isinstance(amount, Decimal) else Decimal(str(amount))
        share = (
            Decimal("0.00")
            if total == 0
            else (value / total * Decimal("100")).quantize(Decimal("0.01"))
        )
        enriched = dict(row)
        enriched["share_percent"] = share
        result.append(enriched)
    return result


def aggregate_daily_operations(purchases: Sequence[Purchase]) -> list[DailyOperationPoint]:
    totals: dict[date, Decimal] = {}
    counts: dict[date, int] = {}
    for purchase in purchases:
        if purchase.date is None:
            continue
        totals[purchase.date] = totals.get(purchase.date, Decimal("0.00")) + purchase.amount
        counts[purchase.date] = counts.get(purchase.date, 0) + 1
    return [
        DailyOperationPoint(
            date=operation_date,
            total_amount=totals[operation_date],
            purchase_count=counts[operation_date],
        )
        for operation_date in sorted(totals)
    ]


def _warnings(
    purchases: Sequence[Purchase],
    participants: Sequence[str],
    scope: str,
) -> list[dict[str, object]]:
    warnings: list[dict[str, object]] = []
    known = set(participants)
    if not purchases:
        warnings.append(
            {
                "warning_type": "no_purchases",
                "purchase_id": "",
                "purchase_name": "",
                "message": (
                    EMPTY_OPEN_SCOPE_MESSAGE if scope == "open" else "Нет покупок для отчета."
                ),
            }
        )
    for purchase in purchases:
        if not purchase.category:
            warnings.append(
                _purchase_warning(
                    "missing_category",
                    purchase,
                    "У покупки не указана категория.",
                )
            )
        if purchase.date is None:
            warnings.append(
                _purchase_warning("missing_date", purchase, "У покупки не указана дата.")
            )
        if not purchase.participants:
            warnings.append(
                _purchase_warning(
                    "empty_participants",
                    purchase,
                    "У покупки пустой список участников.",
                )
            )
        names = [purchase.payer, *purchase.participants]
        for name in sorted({item for item in names if item and item not in known}):
            warnings.append(
                _purchase_warning(
                    "unknown_participant",
                    purchase,
                    f"Участник не найден в справочнике: {name}",
                )
            )
    return warnings


def _purchase_warning(warning_type: str, purchase: Purchase, message: str) -> dict[str, object]:
    return {
        "warning_type": warning_type,
        "purchase_id": purchase.id,
        "purchase_name": purchase.purchase_name or DEFAULT_PURCHASE_NAME,
        "message": message,
    }


def _summary_rows(dataset: CurrentReportDataset) -> list[list[object]]:
    summary = dataset.summary
    return [
        [label_for_report_field("Scope"), label_for_scope(str(summary["scope"]))],
        [label_for_report_field("Snapshot"), summary["generated_at"]],
        ["Общая сумма", summary["total_amount"]],
        ["Количество покупок", summary["purchase_count"]],
        ["Средний чек", summary["average_purchase"]],
        ["Количество участников", summary["participant_count"]],
        ["Итоговые переводы", summary["settlement_count"]],
    ]


def _purchase_rows(dataset: CurrentReportDataset) -> Iterable[list[object]]:
    for row in sorted_purchases_with_payer_totals(dataset.purchases):
        purchase = row.purchase
        yield [
            purchase.id,
            purchase.date,
            purchase.purchase_name or DEFAULT_PURCHASE_NAME,
            purchase.category or UNCATEGORIZED,
            purchase.amount,
            purchase.payer,
            purchase.participants,
            label_for_purchase_status("settled" if purchase.settled else "open"),
            purchase.comment,
            row.payer_total,
            row.payer_rank,
        ]


def _warning_rows(warnings: Sequence[dict[str, object]]) -> Iterable[list[object]]:
    for row in warnings:
        yield [
            row.get("warning_type", ""),
            row.get("purchase_id", ""),
            row.get("purchase_name", ""),
            row.get("message", ""),
        ]


def _write_csv(path: Path, headers: Sequence[str], rows: Iterable[Sequence[object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    materialized_rows = [list(row) for row in rows]
    output_headers = _headers_for_rows(path, headers, materialized_rows)
    with path.open("w", encoding=CSV_ENCODING, newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(output_headers)
        writer.writerows([_serialize(value) for value in row] for row in materialized_rows)


def _headers_for_rows(
    path: Path,
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
) -> list[str]:
    result = list(headers)
    max_columns = max((len(row) for row in rows), default=len(result))
    if path.name == "purchases.csv" and max_columns == len(result) + 2:
        result.extend(["payer_total", "payer_rank"])
    while len(result) < max_columns:
        result.append(f"extra_{len(result) + 1}")
    return result


def _read_csv(path: Path) -> list[list[str]]:
    with path.open(encoding=CSV_ENCODING, newline="") as stream:
        return list(csv.reader(stream))


def _serialize(value: object) -> str:
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (list, tuple)):
        return ", ".join(str(item) for item in value)
    if value is None:
        return ""
    return str(value)


def _json_value(value: object) -> object:
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    return value


def _participant_names(participants: Sequence[Participant | str]) -> list[str]:
    return [
        participant.name if isinstance(participant, Participant) else str(participant)
        for participant in participants
    ]


def _append_markdown_table(
    lines: list[str],
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
) -> None:
    serialized_rows = [[_escape_markdown(value) for value in row] for row in rows]
    output_headers = _headers_for_markdown(headers, serialized_rows)
    lines.extend([f"## {title}", ""])
    if not serialized_rows:
        empty_message = "Предупреждений нет." if title == "Предупреждения" else "Нет данных."
        lines.extend([empty_message, ""])
        return
    lines.append("| " + " | ".join(output_headers) + " |")
    lines.append("|" + "|".join("---" for _ in output_headers) + "|")
    lines.extend("| " + " | ".join(row) + " |" for row in serialized_rows)
    lines.append("")


def _headers_for_markdown(headers: Sequence[str], rows: Sequence[Sequence[str]]) -> list[str]:
    result = list(headers)
    max_columns = max((len(row) for row in rows), default=len(result))
    if max_columns == len(result) + 2:
        result.extend(["payer_total", "payer_rank"])
    while len(result) < max_columns:
        result.append(f"extra_{len(result) + 1}")
    return result


def _escape_markdown(value: object) -> str:
    return _serialize(value).replace("|", "\\|").replace("\n", " ")


def _html_table(report_dir: Path, filename: str, title: str) -> str:
    path = report_dir / "tables" / filename
    if not path.exists():
        return (
            f'<article class="panel"><h3>{escape(title)}</h3>'
            '<p class="missing">Таблица не создана.</p></article>'
        )
    rows = rows_for_visual_table(filename, _read_csv(path))
    if not rows:
        body = '<p class="missing">Нет данных.</p>'
    else:
        header = "".join(f"<th>{escape(cell)}</th>" for cell in rows[0])
        data = "".join(
            "<tr>" + "".join(_html_cell(cell) for cell in row) + "</tr>"
            for row in rows[1:]
        )
        body = f"<table><thead><tr>{header}</tr></thead><tbody>{data}</tbody></table>"
    note = visual_table_note(filename)
    if note:
        body += f'<p class="muted">{escape(note)}</p>'
    return f'<article class="panel"><h3>{escape(title)}</h3>{body}</article>'


def _html_cell(value: str) -> str:
    css_class = ' class="num"' if _is_numeric_text(value) else ""
    return f"<td{css_class}>{escape(value)}</td>"


def _is_numeric_text(value: str) -> bool:
    try:
        Decimal(value)
    except (InvalidOperation, ValueError):
        return False
    return True


def _html_chart(
    path_or_dir: Path,
    filename: str | None = None,
    title: str | None = None,
) -> str:
    if filename is not None:
        path = path_or_dir / "charts" / filename
        chart_title = title or chart_display_title(filename)
        if not path.exists():
            return (
                f'<article class="panel"><h3>{escape(chart_title)}</h3>'
                '<p class="missing">График не создан: нет данных.</p></article>'
            )
    else:
        path = path_or_dir
        chart_title = chart_display_title(path)
    return (
        f'<article class="panel"><h3>{escape(chart_title)}</h3>'
        f'<img class="chart" src="charts/{escape(path.name)}" alt="{escape(chart_title)}">'
        "</article>"
    )


def _chart_sort_key(path: Path) -> int:
    order = {filename: index for index, (filename, _title) in enumerate(CHART_SPECS)}
    return order.get(path.name, len(order))


def _save_barh(
    path: Path,
    labels: list[str],
    values: list[float],
    colors: list[str],
    title: str,
    xlabel: str,
    *,
    non_negative_x_axis: bool = False,
    value_label_colors: list[str] | None = None,
) -> None:
    height = max(4.0, min(9.0, 1.0 + len(labels) * 0.5))
    fig, ax = plt.subplots(figsize=(10, height))
    positions = range(len(labels))
    bars = ax.barh(positions, values, color=colors)
    ax.set_yticks(list(positions), labels=labels)
    ax.invert_yaxis()
    ax.set_title(title)
    ax.set_xlabel(xlabel)
    ax.grid(axis="x", alpha=0.2)
    if non_negative_x_axis:
        set_non_negative_x_axis(ax, values)
    else:
        _pad_axis(ax, values)
    value_labels = ax.bar_label(
        bars,
        labels=[_format_money(value) for value in values],
        padding=6,
        fontsize=9,
    )
    if value_label_colors is not None:
        for label, color in zip(value_labels, value_label_colors):
            label.set_color(color)
    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)


def _pad_axis(ax, values: Sequence[float]) -> None:
    low = min(0.0, min(values))
    high = max(0.0, max(values))
    span = high - low or max(abs(high), abs(low), 1.0)
    pad = span * 0.24
    ax.set_xlim(low - pad, high + pad)


def set_non_negative_x_axis(ax, values: Sequence[float]) -> None:
    if not values:
        return
    if any(value < 0 for value in values):
        _pad_axis(ax, values)
        return
    high = max(values)
    right = high * 1.18 if high > 0 else 1.0
    ax.set_xlim(left=0, right=right)


def _format_money(value: float) -> str:
    return f"{value:,.2f}".replace(",", " ")


def _line_label_offset(index: int, values: Sequence[float]) -> tuple[int, int, str]:
    value = values[index]
    previous_value = values[index - 1] if index > 0 else None
    next_value = values[index + 1] if index < len(values) - 1 else None

    if previous_value is not None and next_value is not None:
        if value <= previous_value and value <= next_value:
            return 0, -14, "top"
        if value >= previous_value and value >= next_value:
            return 0, 10, "bottom"

    neighbor = next_value if previous_value is None else previous_value
    if neighbor is not None and value < neighbor:
        return 0, -14, "top"
    return 0, 10, "bottom"


def _chart_category(dataset: CurrentReportDataset, path: Path) -> None:
    labels = [str(row["category"]) for row in dataset.by_category]
    colors = build_stable_color_map(sorted(labels))
    _save_barh(
        path,
        labels,
        [float(row["total_amount"]) for row in dataset.by_category],
        [colors[label] for label in labels],
        "Расходы по категориям",
        "Сумма",
        non_negative_x_axis=True,
    )


def _chart_payer(dataset: CurrentReportDataset, path: Path) -> None:
    source_labels = [str(row["payer"]) for row in dataset.by_payer]
    labels = [
        truncate_display_label(label, PARTICIPANT_LABEL_MAX_LENGTH, "Без имени")
        for label in source_labels
    ]
    colors = build_stable_color_map(source_labels)
    _save_barh(
        path,
        labels,
        [float(row["total_paid"]) for row in dataset.by_payer],
        [colors[label] for label in source_labels],
        "Расходы по плательщикам",
        "Сумма",
        non_negative_x_axis=True,
    )


def _chart_participant_share(dataset: CurrentReportDataset, path: Path) -> None:
    source_labels = [str(row["participant"]) for row in dataset.by_participant]
    labels = [
        truncate_display_label(label, PARTICIPANT_LABEL_MAX_LENGTH, "Без имени")
        for label in source_labels
    ]
    colors = build_stable_color_map(source_labels)
    _save_barh(
        path,
        labels,
        [float(row["total_share"]) for row in dataset.by_participant],
        [colors[label] for label in source_labels],
        chart_display_title("participant_share.png"),
        "Сумма",
        non_negative_x_axis=True,
    )


def _chart_balances(dataset: CurrentReportDataset, path: Path) -> None:
    rows = sorted(dataset.balances, key=lambda row: row.net)
    colors = [color_for_balance_status(row.net) for row in rows]
    _save_barh(
        path,
        [
            truncate_display_label(
                row.participant,
                PARTICIPANT_LABEL_MAX_LENGTH,
                "Без имени",
            )
            for row in rows
        ],
        [float(row.net) for row in rows],
        colors,
        BALANCE_CHART_TITLE,
        BALANCE_CHART_XLABEL,
        value_label_colors=colors,
    )


def _chart_operations_by_day(dataset: CurrentReportDataset, path: Path) -> None:
    rows = aggregate_daily_operations(dataset.purchases)
    labels = [row.date.isoformat() for row in rows]
    values = [float(row.total_amount) for row in rows]
    period_colors = build_period_color_map(labels)
    fig, ax = plt.subplots(figsize=(10, 5))
    ax.plot(labels, values, color=EXPENSE_DIMENSION_COLOR_MAP["Период"])
    ax.scatter(labels, values, color=[period_colors[label] for label in labels], zorder=3)
    for index, (label, value) in enumerate(zip(labels, values)):
        x_offset, y_offset, vertical_alignment = _line_label_offset(index, values)
        ax.annotate(
            _format_money(value),
            (label, value),
            textcoords="offset points",
            xytext=(x_offset, y_offset),
            ha="center",
            va=vertical_alignment,
            fontsize=9,
        )
    ax.set_title(chart_display_title(OPERATIONS_BY_DAY_FILENAME))
    ax.set_xlabel("Дата")
    ax.set_ylabel("Сумма")
    ax.grid(alpha=0.2)
    high = max(values) if values else 0.0
    ax.set_ylim(bottom=0, top=high * 1.18 if high > 0 else 1.0)
    ax.tick_params(axis="x", rotation=45)
    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)


def _chart_top_purchases(dataset: CurrentReportDataset, path: Path) -> None:
    rows = dataset.top_purchases
    _save_barh(
        path,
        [
            truncate_display_label(
                row["purchase_name"],
                PURCHASE_LABEL_MAX_LENGTH,
                "н/д",
            )
            for row in rows
        ],
        [float(row["amount"]) for row in rows],
        [QUALITATIVE_PALETTE[(int(row["rank"]) - 1) % len(QUALITATIVE_PALETTE)] for row in rows],
        "Крупнейшие покупки",
        "Сумма",
        non_negative_x_axis=True,
    )


def _chart_warning(filename: str) -> dict[str, object]:
    return {
        "chart": Path(filename).stem,
        "warning_type": "chart_no_data",
        "purchase_id": "",
        "purchase_name": "",
        "message": (
            f"График «{chart_display_title(filename)}» не создан: "
            "нет данных для текущего режима расчета."
        ),
    }


def _operations_by_day_not_enough_data_warning() -> dict[str, object]:
    return {
        "chart": Path(OPERATIONS_BY_DAY_FILENAME).stem,
        "warning_type": OPERATIONS_BY_DAY_WARNING_TYPE,
        "purchase_id": "",
        "purchase_name": "",
        "message": OPERATIONS_BY_DAY_WARNING_MESSAGE,
    }


def _write_sheet(worksheet, sheet_name: str, rows: list[list[str]]) -> None:
    primary = "285F8F"
    header = "EAF1F7"
    ink = "18212F"
    thin = Side(style="thin", color="DCE3EC")
    worksheet.sheet_view.showGridLines = False
    max_columns = max((len(row) for row in rows), default=2)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_columns)
    title = worksheet.cell(1, 1, sheet_name)
    title.font = Font(size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor=primary)
    title.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 28

    if not rows:
        worksheet.cell(3, 1, "Нет данных.")
        return

    headers = rows[0]
    for column, value in enumerate(headers, 1):
        cell = worksheet.cell(3, column, value)
        cell.font = Font(bold=True, color=ink)
        cell.fill = PatternFill("solid", fgColor=header)
        cell.border = Border(bottom=thin)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row_index, row in enumerate(rows[1:], 4):
        for column, raw_value in enumerate(row, 1):
            value = _typed_value(raw_value, headers[column - 1], row)
            cell = worksheet.cell(row_index, column, value)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if _is_money(headers[column - 1], row):
                cell.number_format = "#,##0.00"
    last_row = max(3, len(rows) + 2)
    last_column = get_column_letter(max_columns)
    worksheet.freeze_panes = "A4"
    worksheet.auto_filter.ref = f"A3:{last_column}{last_row}"
    for row_index in range(4, last_row + 1):
        if row_index % 2 == 0:
            for cell in worksheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
    for index, value in enumerate(headers, 1):
        column_values = [value] + [row[index - 1] for row in rows[1:] if index <= len(row)]
        width = min(max(max((len(str(item)) for item in column_values), default=8) + 2, 12), 42)
        worksheet.column_dimensions[get_column_letter(index)].width = width


def _typed_value(raw_value: str, header: str, row: list[str]):
    if raw_value == "":
        return None
    if header == "payer_rank":
        try:
            return int(raw_value)
        except ValueError:
            return raw_value
    if header == "Дата":
        try:
            return date.fromisoformat(raw_value)
        except ValueError:
            return raw_value
    if _is_money(header, row):
        try:
            return float(raw_value)
        except ValueError:
            return raw_value
    if header in {"Покупок", "Количество покупок", "Количество участников", "Итоговые переводы"}:
        try:
            return int(raw_value)
        except ValueError:
            return raw_value
    return raw_value


def _is_money(header: str, row: list[str]) -> bool:
    if header == "payer_total":
        return True
    money_headers = {
        "Сумма",
        "Оплачено",
        "Доля расходов",
        "Доля",
        "Баланс",
        "Объем расходов на человека",
        BALANCE_RESULT_HEADER,
        "Доля, %",
        "Доля оплат, %",
    }
    money_summary = {"Общая сумма", "Средний чек"}
    return header in money_headers or (header == "Значение" and row and row[0] in money_summary)


def _write_charts_sheet(
    worksheet,
    charts_dir: Path,
    chart_paths: Sequence[Path] = (),
) -> None:
    primary = "285F8F"
    muted = "64748B"
    secondary = "DCE8F2"
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:J1")
    worksheet["A1"] = "Графики"
    worksheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor=primary)
    worksheet.row_dimensions[1].height = 28
    current_chart_paths = (
        sorted(chart_paths, key=_chart_sort_key)
        if chart_paths
        else sorted(charts_dir.glob("*.png"))
        if charts_dir.exists()
        else []
    )
    if not current_chart_paths:
        worksheet["A3"] = "Графики не созданы: нет данных."
        worksheet["A3"].font = Font(italic=True, color=muted)
        worksheet["A3"].fill = PatternFill("solid", fgColor=secondary)
        return
    row = 3
    for chart_path in current_chart_paths:
        worksheet.cell(row, 1, chart_display_title(chart_path))
        worksheet.cell(row, 1).font = Font(bold=True, color=primary)
        row += 1
        image = Image(chart_path)
        if image.width > 900:
            ratio = 900 / image.width
            image.width = 900
            image.height = int(image.height * ratio)
        worksheet.add_image(image, f"A{row}")
        row += max(24, int(image.height / 20) + 3)
    for column in range(1, 11):
        worksheet.column_dimensions[get_column_letter(column)].width = 12
