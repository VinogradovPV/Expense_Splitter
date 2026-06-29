from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal
from html import escape
from pathlib import Path
from typing import Iterable, Sequence

from expense_splitter.analytics import (
    UNCATEGORIZED,
    aggregate_by_category,
    aggregate_by_participant,
    aggregate_by_payer,
    build_top_purchases,
)
from expense_splitter.calculator import calculate_balances
from expense_splitter.current_report import (
    CHART_SPECS,
    CURRENT_HTML_CSS,
    SUPPORTED_FORMATS,
    _append_markdown_table,
    _html_chart,
    _html_table,
    _json_value,
    _read_csv,
    _serialize,
    _with_share_percent,
    _write_charts_sheet,
    _write_csv,
    _write_sheet,
)
from expense_splitter.models import (
    DEFAULT_PURCHASE_NAME,
    Balance,
    Participant,
    Purchase,
    Settlement,
    SettlementPeriod,
)
from expense_splitter.settlement_periods import get_settlement_period
from expense_splitter.visual.palette import PALETTE_NAME, PALETTE_VERSION

CSV_ENCODING = "utf-8-sig"
REPORT_TITLE = "Отчет по периоду взаиморасчетов"
BALANCE_SOURCE_RECOMPUTED = "recomputed_from_period_purchases"
SETTLEMENT_SOURCE_SNAPSHOT = "snapshot"

TABLE_SPECS = (
    ("summary.csv", "Сводка"),
    ("purchases.csv", "Покупки периода"),
    ("by_category.csv", "Расходы по категориям"),
    ("by_payer.csv", "Расходы по плательщикам"),
    ("by_participant.csv", "Доли участников"),
    ("balances.csv", "Балансы"),
    ("settlements.csv", "Итоговые переводы"),
    ("warnings.csv", "Warnings"),
)


class SettlementPeriodReportError(ValueError):
    """User-facing settlement period report error."""


@dataclass(frozen=True)
class SettlementPeriodReportDataset:
    period: SettlementPeriod
    generated_at: datetime
    participants: list[str]
    purchases: list[Purchase]
    missing_purchase_ids: list[str]
    balances: list[Balance]
    settlements: list[Settlement]
    summary: dict[str, object]
    by_category: list[dict[str, object]]
    by_payer: list[dict[str, object]]
    by_participant: list[dict[str, object]]
    top_purchases: list[dict[str, object]]
    warnings: list[dict[str, object]]
    balance_source: str
    settlement_source: str


def build_settlement_period_report_dataset(
    periods: Sequence[SettlementPeriod],
    purchases: Sequence[Purchase],
    participants: Sequence[Participant | str],
    settlement_period_id: str,
    generated_at: datetime | None = None,
) -> SettlementPeriodReportDataset:
    try:
        period = get_settlement_period(periods, settlement_period_id)
    except ValueError as exc:
        raise SettlementPeriodReportError(
            f"Период взаиморасчетов не найден: {settlement_period_id}"
        ) from exc

    purchase_by_id = {purchase.id: purchase for purchase in purchases}
    selected: list[Purchase] = []
    missing: list[str] = []
    for purchase_id in period.purchase_ids:
        purchase = purchase_by_id.get(purchase_id)
        if purchase is None:
            missing.append(purchase_id)
        else:
            selected.append(purchase)

    participant_names = _period_participants(participants, selected, period.settlements)
    balances = calculate_balances(selected, participant_names) if selected else []
    generated = generated_at or datetime.now(timezone.utc).astimezone()
    warnings = _warnings(period, selected, missing)
    total_amount = period.total_amount
    summary = _summary(period, selected, participant_names, generated)

    return SettlementPeriodReportDataset(
        period=period,
        generated_at=generated,
        participants=participant_names,
        purchases=selected,
        missing_purchase_ids=missing,
        balances=balances,
        settlements=list(period.settlements),
        summary=summary,
        by_category=_with_share_percent(aggregate_by_category(selected), total_amount),
        by_payer=aggregate_by_payer(selected),
        by_participant=aggregate_by_participant(selected, participant_names) if selected else [],
        top_purchases=build_top_purchases(selected) if selected else [],
        warnings=warnings,
        balance_source=BALANCE_SOURCE_RECOMPUTED,
        settlement_source=SETTLEMENT_SOURCE_SNAPSHOT,
    )


def generate_settlement_period_report(
    dataset: SettlementPeriodReportDataset,
    output_root: Path = Path("reports/settlement_periods"),
    output_format: str = "all",
) -> Path:
    format_key = output_format.lower()
    if format_key not in SUPPORTED_FORMATS:
        raise ValueError("Format must be one of: markdown, csv, png, html, xlsx, all.")

    report_dir = output_root / f"{dataset.period.id}_{dataset.generated_at:%Y-%m-%d_%H-%M-%S}"
    tables_dir = report_dir / "tables"
    charts_dir = report_dir / "charts"
    tables_dir.mkdir(parents=True, exist_ok=True)
    charts_dir.mkdir(parents=True, exist_ok=True)

    warnings = [dict(row) for row in dataset.warnings]
    generated: list[Path] = []

    if format_key in {"csv", "html", "xlsx", "all"}:
        generated.extend(write_settlement_period_csv_tables(dataset, tables_dir, warnings))
    else:
        generated.extend(write_settlement_period_minimum_tables(dataset, tables_dir, warnings))

    if format_key in {"png", "html", "xlsx", "all"}:
        chart_paths, chart_warnings = write_settlement_period_charts(dataset, charts_dir)
        generated.extend(chart_paths)
        warnings.extend(chart_warnings)
        write_settlement_period_warnings_table(tables_dir / "warnings.csv", warnings)

    if format_key in {"markdown", "html", "all"}:
        generated.append(
            write_settlement_period_markdown(
                dataset,
                report_dir / "settlement_period_report.md",
                warnings,
            )
        )

    if format_key in {"html", "all"}:
        generated.append(
            write_settlement_period_html(
                dataset,
                report_dir / "settlement_period_dashboard.html",
                warnings,
            )
        )

    if format_key in {"xlsx", "all"}:
        generated.append(
            write_settlement_period_xlsx(
                report_dir / "settlement_period.xlsx",
                tables_dir,
                charts_dir,
            )
        )

    metadata_path = report_dir / "metadata.json"
    generated.append(metadata_path)
    write_settlement_period_metadata(dataset, metadata_path, generated, warnings, report_dir)
    return report_dir


def write_settlement_period_csv_tables(
    dataset: SettlementPeriodReportDataset,
    tables_dir: Path,
    warnings: Sequence[dict[str, object]] | None = None,
) -> list[Path]:
    warning_rows = list(dataset.warnings if warnings is None else warnings)
    specs = (
        ("summary.csv", ["Показатель", "Значение"], _summary_rows(dataset)),
        (
            "purchases.csv",
            [
                "ID",
                "Дата",
                "Покупка",
                "Категория",
                "Сумма",
                "Плательщик",
                "Участники",
                "Статус",
                "Комментарий",
            ],
            _purchase_rows(dataset),
        ),
        (
            "by_category.csv",
            ["Категория", "Сумма", "Покупок", "Доля, %"],
            (
                [
                    row["category"],
                    row["total_amount"],
                    row["purchase_count"],
                    row["share_percent"],
                ]
                for row in dataset.by_category
            ),
        ),
        (
            "by_payer.csv",
            ["Плательщик", "Оплачено", "Покупок"],
            ([row["payer"], row["total_paid"], row["purchase_count"]] for row in dataset.by_payer),
        ),
        (
            "by_participant.csv",
            ["Участник", "Доля расходов", "Покупок"],
            (
                [row["participant"], row["total_share"], row["purchase_count"]]
                for row in dataset.by_participant
            ),
        ),
        (
            "balances.csv",
            ["Участник", "Оплачено", "Доля", "Баланс"],
            ([row.participant, row.paid, row.share, row.net] for row in dataset.balances),
        ),
        (
            "settlements.csv",
            ["От кого", "Кому", "Сумма"],
            ([row.from_participant, row.to_participant, row.amount] for row in dataset.settlements),
        ),
        (
            "warnings.csv",
            ["Тип", "ID покупки", "Покупка", "Сообщение"],
            _warning_rows(warning_rows),
        ),
    )
    paths: list[Path] = []
    for filename, headers, rows in specs:
        path = tables_dir / filename
        _write_csv(path, headers, rows)
        paths.append(path)
    return paths


def write_settlement_period_minimum_tables(
    dataset: SettlementPeriodReportDataset,
    tables_dir: Path,
    warnings: Sequence[dict[str, object]],
) -> list[Path]:
    return [
        write_settlement_period_warnings_table(tables_dir / "warnings.csv", warnings),
        _write_summary_table(tables_dir / "summary.csv", dataset),
    ]


def write_settlement_period_warnings_table(
    path: Path,
    warnings: Sequence[dict[str, object]],
) -> Path:
    _write_csv(path, ["Тип", "ID покупки", "Покупка", "Сообщение"], _warning_rows(warnings))
    return path


def write_settlement_period_markdown(
    dataset: SettlementPeriodReportDataset,
    path: Path,
    warnings: Sequence[dict[str, object]],
) -> Path:
    period = dataset.period
    summary = dataset.summary
    lines = [
        f"# {REPORT_TITLE}",
        "",
        f"Settlement period ID: {period.id}",
        f"Название: {period.name}",
        f"Статус: {period.status}",
        f"Даты: {period.date_from.isoformat()} — {period.date_to.isoformat()}",
        f"Closed at: {period.closed_at.isoformat(timespec='seconds')}",
        f"Reopened at: {_serialize(period.reopened_at)}",
        f"Snapshot: {dataset.generated_at.isoformat(timespec='seconds')}",
        "",
        "## Сводка",
        "",
        "| Показатель | Значение |",
        "|---|---:|",
        f"| Общая сумма периода | {_serialize(summary['total_amount'])} |",
        f"| Количество покупок | {summary['purchase_count']} |",
        f"| Средний чек | {_serialize(summary['average_purchase'])} |",
        f"| Количество участников | {summary['participant_count']} |",
        f"| Количество итоговых переводов | {summary['settlement_count']} |",
        f"| Статус периода | {summary['status']} |",
        "",
    ]
    if not dataset.purchases:
        lines.extend(["Период не содержит покупок.", ""])

    _append_markdown_table(
        lines,
        "Расходы по категориям",
        ["Категория", "Сумма", "Покупок", "Доля, %"],
        (
            [row["category"], row["total_amount"], row["purchase_count"], row["share_percent"]]
            for row in dataset.by_category
        ),
    )
    _append_markdown_table(
        lines,
        "Расходы по плательщикам",
        ["Плательщик", "Оплачено", "Покупок"],
        ([row["payer"], row["total_paid"], row["purchase_count"]] for row in dataset.by_payer),
    )
    _append_markdown_table(
        lines,
        "Доли участников",
        ["Участник", "Доля расходов", "Покупок"],
        (
            [row["participant"], row["total_share"], row["purchase_count"]]
            for row in dataset.by_participant
        ),
    )
    _append_markdown_table(
        lines,
        "Балансы",
        ["Участник", "Оплачено", "Доля", "Баланс"],
        ([row.participant, row.paid, row.share, row.net] for row in dataset.balances),
    )
    _append_markdown_table(
        lines,
        "Итоговые переводы",
        ["От кого", "Кому", "Сумма"],
        ([row.from_participant, row.to_participant, row.amount] for row in dataset.settlements),
    )
    _append_markdown_table(
        lines,
        "Покупки периода",
        [
            "Дата",
            "Покупка",
            "Категория",
            "Сумма",
            "Плательщик",
            "Участники",
            "Статус",
            "Комментарий",
        ],
        (
            [row[1], row[2], row[3], row[4], row[5], row[6], row[7], row[8]]
            for row in _purchase_rows(dataset)
        ),
    )
    _append_markdown_table(
        lines,
        "Warnings",
        ["Тип", "ID покупки", "Покупка", "Сообщение"],
        _warning_rows(warnings),
    )

    chart_dir = path.parent / "charts"
    chart_paths = sorted(chart_dir.glob("*.png")) if chart_dir.exists() else []
    if chart_paths:
        lines.extend(["## Графики", ""])
        for chart_path in chart_paths:
            lines.extend([f"![{chart_path.stem}](charts/{chart_path.name})", ""])

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_settlement_period_html(
    dataset: SettlementPeriodReportDataset,
    path: Path,
    warnings: Sequence[dict[str, object]],
) -> Path:
    period = dataset.period
    summary = dataset.summary
    cards = (
        ("Общая сумма", _serialize(summary["total_amount"])),
        ("Покупок", str(summary["purchase_count"])),
        ("Средний чек", _serialize(summary["average_purchase"])),
        ("Участников", str(summary["participant_count"])),
        ("Переводов", str(summary["settlement_count"])),
        ("Статус", str(summary["status"])),
    )
    card_html = "".join(
        f'<div class="card"><span>{escape(label)}</span><strong>{escape(value)}</strong></div>'
        for label, value in cards
    )
    table_html = "".join(
        _html_table(path.parent, filename, title) for filename, title in TABLE_SPECS
    )
    chart_html = "".join(
        _html_chart(path.parent, filename, title) for filename, title in CHART_SPECS
    )
    warning_items = "".join(f"<li>{escape(str(row.get('message', '')))}</li>" for row in warnings)
    warnings_html = f"<ul>{warning_items}</ul>" if warning_items else "<p>Предупреждений нет.</p>"
    csv_links = "".join(
        f'<a href="tables/{escape(filename)}">{escape(title)} CSV</a>'
        for filename, title in TABLE_SPECS
    )
    empty_note = (
        '<p class="empty">Период не содержит покупок.</p>' if not dataset.purchases else ""
    )
    reopened_at = _serialize(period.reopened_at)
    document = f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{escape(REPORT_TITLE)}</title>
<style>{CURRENT_HTML_CSS}</style></head><body><main>
<header><h1>{escape(REPORT_TITLE)}</h1>
<p class="meta">ID: {escape(period.id)}<br>
Название: {escape(period.name)}<br>
Статус: {escape(period.status)}<br>
Даты: {escape(period.date_from.isoformat())} — {escape(period.date_to.isoformat())}<br>
Closed at: {escape(period.closed_at.isoformat(timespec='seconds'))}<br>
Reopened at: {escape(reopened_at)}<br>
Snapshot: {escape(dataset.generated_at.isoformat(timespec='seconds'))}</p></header>
<section><h2>Сводка</h2>{empty_note}<div class="cards">{card_html}</div></section>
<section><h2>Warnings</h2><div class="panel warning">{warnings_html}</div></section>
<section><h2>Графики</h2>{chart_html}</section>
<section><h2>Таблицы</h2>{table_html}</section>
<section><h2>Файлы</h2><div class="panel links">
<a href="settlement_period_report.md">Markdown</a>{csv_links}</div></section>
</main></body></html>"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(document, encoding="utf-8")
    return path


def write_settlement_period_xlsx(path: Path, tables_dir: Path, charts_dir: Path) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = (
        ("Summary", "summary.csv"),
        ("Purchases", "purchases.csv"),
        ("By Category", "by_category.csv"),
        ("By Payer", "by_payer.csv"),
        ("By Participant", "by_participant.csv"),
        ("Balances", "balances.csv"),
        ("Settlements", "settlements.csv"),
        ("Warnings", "warnings.csv"),
    )
    for sheet_name, filename in sheets:
        worksheet = workbook.create_sheet(sheet_name)
        _write_sheet(worksheet, sheet_name, _read_csv(tables_dir / filename))
    charts_sheet = workbook.create_sheet("Charts")
    _write_charts_sheet(charts_sheet, charts_dir)
    workbook.active = 0
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def write_settlement_period_charts(
    dataset: SettlementPeriodReportDataset,
    output_dir: Path,
) -> tuple[list[Path], list[dict[str, object]]]:
    from expense_splitter.current_report import (
        _chart_balances,
        _chart_category,
        _chart_participant_share,
        _chart_payer,
        _chart_top_purchases,
    )

    output_dir.mkdir(parents=True, exist_ok=True)
    generated: list[Path] = []
    warnings: list[dict[str, object]] = []
    specs = (
        ("spending_by_category.png", dataset.by_category, _chart_category),
        ("spending_by_payer.png", dataset.by_payer, _chart_payer),
        ("participant_share.png", dataset.by_participant, _chart_participant_share),
        ("balances.png", dataset.balances, _chart_balances),
        ("top_purchases.png", dataset.top_purchases, _chart_top_purchases),
    )
    for filename, rows, renderer in specs:
        if not dataset.purchases or not rows:
            warnings.append(_chart_warning(filename))
            continue
        path = output_dir / filename
        renderer(dataset, path)  # type: ignore[arg-type]
        generated.append(path)
    return generated, warnings


def write_settlement_period_metadata(
    dataset: SettlementPeriodReportDataset,
    path: Path,
    generated_files: Sequence[Path],
    warnings: Sequence[dict[str, object]],
    report_dir: Path,
) -> Path:
    period = dataset.period
    metadata = {
        "schema_version": 1,
        "report_type": "settlement_period",
        "settlement_period_id": period.id,
        "period_status": period.status,
        "generated_at": dataset.generated_at.isoformat(),
        "snapshot_source": {
            "purchase_ids": "settlement_periods.yaml",
            "total_amount": "settlement_periods.yaml",
            "settlements": dataset.settlement_source,
            "balance_source": dataset.balance_source,
            "purchase_details": "current_purchases_by_snapshot_ids",
        },
        "summary": {key: _metadata_value(value) for key, value in dataset.summary.items()},
        "missing_purchase_ids": list(dataset.missing_purchase_ids),
        "palette": {"name": PALETTE_NAME, "version": PALETTE_VERSION},
        "warning_count": len(warnings),
        "files": sorted(
            str(item.relative_to(report_dir)).replace("\\", "/")
            for item in generated_files
            if item.exists()
        ),
    }
    path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def _metadata_value(value: object) -> object:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return _json_value(value)


def _period_participants(
    participants: Sequence[Participant | str],
    purchases: Sequence[Purchase],
    settlements: Sequence[Settlement],
) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for participant in participants:
        name = participant.name if isinstance(participant, Participant) else str(participant)
        _append_unique_name(result, seen, name)
    for purchase in purchases:
        _append_unique_name(result, seen, purchase.payer)
        for participant in purchase.participants:
            _append_unique_name(result, seen, participant)
    for settlement in settlements:
        _append_unique_name(result, seen, settlement.from_participant)
        _append_unique_name(result, seen, settlement.to_participant)
    return result


def _append_unique_name(result: list[str], seen: set[str], name: str) -> None:
    value = str(name).strip()
    key = value.casefold()
    if value and key not in seen:
        result.append(value)
        seen.add(key)


def _summary(
    period: SettlementPeriod,
    purchases: Sequence[Purchase],
    participants: Sequence[str],
    generated_at: datetime,
) -> dict[str, object]:
    purchase_count = len(purchases)
    average_purchase = (
        (period.total_amount / Decimal(purchase_count)).quantize(Decimal("0.01"))
        if purchase_count
        else Decimal("0.00")
    )
    return {
        "settlement_period_id": period.id,
        "name": period.name,
        "status": period.status,
        "date_from": period.date_from,
        "date_to": period.date_to,
        "closed_at": period.closed_at.isoformat(timespec="seconds"),
        "reopened_at": period.reopened_at.isoformat(timespec="seconds")
        if period.reopened_at
        else "",
        "generated_at": generated_at.isoformat(timespec="seconds"),
        "total_amount": period.total_amount,
        "purchase_count": purchase_count,
        "average_purchase": average_purchase,
        "participant_count": len(participants),
        "settlement_count": len(period.settlements),
    }


def _warnings(
    period: SettlementPeriod,
    purchases: Sequence[Purchase],
    missing_purchase_ids: Sequence[str],
) -> list[dict[str, object]]:
    warnings: list[dict[str, object]] = []
    if not period.purchase_ids:
        warnings.append(_warning("no_purchases", "", "", "Период не содержит покупок."))
    if period.status == "reopened":
        warnings.append(
            _warning(
                "period_reopened",
                "",
                "",
                "Период был переоткрыт. Отчет показывает сохраненный snapshot периода.",
            )
        )
    if not period.settlements:
        warnings.append(
            _warning(
                "missing_snapshot_settlements",
                "",
                "",
                "В snapshot периода нет итоговых переводов.",
            )
        )
    for purchase_id in missing_purchase_ids:
        warnings.append(
            _warning(
                "missing_purchase",
                purchase_id,
                "",
                f"purchase_id из snapshot не найден в текущем purchases.yaml: {purchase_id}",
            )
        )
    period_ids = set(period.purchase_ids)
    for purchase in purchases:
        if purchase.id not in period_ids:
            continue
        if purchase.settlement_period_id != period.id:
            warnings.append(
                _purchase_warning(
                    "purchase_current_state_diff",
                    purchase,
                    "Покупка сейчас не привязана к этому settlement period.",
                )
            )
        if not purchase.category:
            warnings.append(
                _purchase_warning(
                    "missing_category",
                    purchase,
                    "У покупки не указана категория.",
                )
            )
    return warnings


def _summary_rows(dataset: SettlementPeriodReportDataset) -> list[list[object]]:
    summary = dataset.summary
    return [
        ["Settlement period ID", summary["settlement_period_id"]],
        ["Название", summary["name"]],
        ["Статус", summary["status"]],
        ["Дата начала", summary["date_from"]],
        ["Дата окончания", summary["date_to"]],
        ["Closed at", summary["closed_at"]],
        ["Reopened at", summary["reopened_at"]],
        ["Snapshot", summary["generated_at"]],
        ["Общая сумма периода", summary["total_amount"]],
        ["Количество покупок", summary["purchase_count"]],
        ["Средний чек", summary["average_purchase"]],
        ["Количество участников", summary["participant_count"]],
        ["Итоговые переводы", summary["settlement_count"]],
        ["Источник балансов", dataset.balance_source],
        ["Источник переводов", dataset.settlement_source],
    ]


def _write_summary_table(path: Path, dataset: SettlementPeriodReportDataset) -> Path:
    _write_csv(path, ["Показатель", "Значение"], _summary_rows(dataset))
    return path


def _purchase_rows(dataset: SettlementPeriodReportDataset) -> Iterable[list[object]]:
    for purchase in sorted(dataset.purchases, key=lambda item: (item.date or date.min, item.id)):
        yield [
            purchase.id,
            purchase.date,
            purchase.purchase_name or DEFAULT_PURCHASE_NAME,
            purchase.category or UNCATEGORIZED,
            purchase.amount,
            purchase.payer,
            purchase.participants,
            "settled" if purchase.settled else "open",
            purchase.comment,
        ]


def _warning_rows(warnings: Sequence[dict[str, object]]) -> Iterable[list[object]]:
    for row in warnings:
        yield [
            row.get("warning_type", ""),
            row.get("purchase_id", ""),
            row.get("purchase_name", ""),
            row.get("message", ""),
        ]


def _warning(
    warning_type: str,
    purchase_id: str,
    purchase_name: str,
    message: str,
) -> dict[str, object]:
    return {
        "warning_type": warning_type,
        "purchase_id": purchase_id,
        "purchase_name": purchase_name,
        "message": message,
    }


def _purchase_warning(warning_type: str, purchase: Purchase, message: str) -> dict[str, object]:
    return _warning(
        warning_type,
        purchase.id,
        purchase.purchase_name or DEFAULT_PURCHASE_NAME,
        message,
    )


def _chart_warning(filename: str) -> dict[str, object]:
    return _warning(
        "chart_no_data",
        "",
        "",
        f"График {filename} не создан: нет данных для settlement period.",
    )
