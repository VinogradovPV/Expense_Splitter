import csv
import json
from datetime import date, datetime, timezone
from decimal import Decimal

from openpyxl import load_workbook
from typer.testing import CliRunner

from expense_splitter.calculator import calculate_balances
from expense_splitter.cli import app
from expense_splitter.current_report import (
    EMPTY_OPEN_SCOPE_MESSAGE,
    build_current_report_dataset,
    generate_current_report,
)
from expense_splitter.models import Participant, Purchase
from expense_splitter.settlement import calculate_settlements
from expense_splitter.storage import initialize_data_files, save_participants, save_purchases

runner = CliRunner()


def sample_participants():
    return [Participant("Alice"), Participant("Bob"), Participant("Cara")]


def sample_purchases():
    return [
        Purchase(
            id="open-1",
            date=date(2026, 6, 1),
            amount=Decimal("90.00"),
            payer="Alice",
            participants=["Alice", "Bob", "Cara"],
            purchase_name="Lunch",
            category="Food",
            comment="shared",
        ),
        Purchase(
            id="open-2",
            date=None,
            amount=Decimal("60.00"),
            payer="Bob",
            participants=["Alice", "Bob"],
            purchase_name="Taxi",
            category=None,
        ),
        Purchase(
            id="settled-1",
            date=date(2026, 6, 2),
            amount=Decimal("30.00"),
            payer="Cara",
            participants=["Alice", "Cara"],
            purchase_name="Coffee",
            category="Food",
            settled=True,
            settlement_period_id="period-1",
        ),
    ]


def test_dataset_open_scope_excludes_settled_purchases_and_reuses_settlement_logic():
    participants = sample_participants()
    purchases = sample_purchases()

    dataset = build_current_report_dataset(participants, purchases, scope="open")

    assert [purchase.id for purchase in dataset.purchases] == ["open-1", "open-2"]
    assert dataset.summary["total_amount"] == Decimal("150.00")
    assert dataset.summary["purchase_count"] == 2
    assert dataset.summary["average_purchase"] == Decimal("75.00")
    assert sum(row["payer_share_percent"] for row in dataset.by_payer) == Decimal("100.00")
    expected_balances = calculate_balances(dataset.purchases, ["Alice", "Bob", "Cara"])
    assert dataset.balances == expected_balances
    assert dataset.settlements == calculate_settlements(expected_balances)


def test_dataset_all_scope_includes_settled_purchases_and_aggregates_rows():
    dataset = build_current_report_dataset(sample_participants(), sample_purchases(), scope="all")

    assert [purchase.id for purchase in dataset.purchases] == ["open-1", "open-2", "settled-1"]
    assert dataset.summary["total_amount"] == Decimal("180.00")
    assert {row["category"] for row in dataset.by_category} == {"Food", "Без категории"}
    food = next(row for row in dataset.by_category if row["category"] == "Food")
    assert food["purchase_count"] == 2
    assert food["share_percent"] == Decimal("66.67")
    assert any(row["warning_type"] == "missing_date" for row in dataset.warnings)
    assert any(row["warning_type"] == "missing_category" for row in dataset.warnings)


def test_generate_current_report_writes_all_outputs(tmp_path):
    dataset = build_current_report_dataset(
        sample_participants(),
        sample_purchases(),
        scope="open",
        generated_at=datetime(2026, 6, 25, 9, 30, 0, tzinfo=timezone.utc),
    )

    report_dir = generate_current_report(dataset, tmp_path, "all")

    assert report_dir == tmp_path / "open_2026-06-25_09-30-00"
    assert (report_dir / "current_state_report.md").is_file()
    assert (report_dir / "current_state_dashboard.html").is_file()
    assert (report_dir / "current_state.xlsx").is_file()
    assert (report_dir / "current_state.pdf").is_file()
    assert (report_dir / "metadata.json").is_file()
    assert {path.name for path in (report_dir / "tables").glob("*.csv")} == {
        "summary.csv",
        "purchases.csv",
        "by_category.csv",
        "by_payer.csv",
        "by_participant.csv",
        "balances.csv",
        "settlements.csv",
        "warnings.csv",
    }
    assert {path.name for path in (report_dir / "charts").glob("*.png")} == {
        "spending_by_category.png",
        "spending_by_payer.png",
        "participant_share.png",
        "balances.png",
        "top_purchases.png",
    }

    metadata = json.loads((report_dir / "metadata.json").read_text(encoding="utf-8"))
    assert metadata["report_type"] == "current_state"
    assert metadata["scope"] == "open"
    assert metadata["summary"]["total_amount"] == "150.00"


def test_current_report_csv_html_markdown_and_xlsx_are_readable(tmp_path):
    dataset = build_current_report_dataset(sample_participants(), sample_purchases(), scope="open")

    report_dir = generate_current_report(dataset, tmp_path, "all")

    purchases_path = report_dir / "tables" / "purchases.csv"
    assert purchases_path.read_bytes().startswith(b"\xef\xbb\xbf")
    with purchases_path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    full_header = rows[0]
    assert full_header[-2:] == ["payer_total", "payer_rank"]
    rows[0] = full_header[:-2]
    assert rows[0] == [
        "ID",
        "Дата",
        "Покупка",
        "Категория",
        "Сумма",
        "Плательщик",
        "Участники",
        "Статус",
        "Комментарий",
    ]
    assert any(row[2] == "Lunch" for row in rows[1:])

    with (report_dir / "tables" / "by_payer.csv").open(
        encoding="utf-8-sig",
        newline="",
    ) as stream:
        payer_rows = list(csv.reader(stream))
    assert payer_rows[0] == ["Плательщик", "Оплачено", "Покупок", "Доля оплат, %"]
    assert Decimal(payer_rows[1][3]) + Decimal(payer_rows[2][3]) == Decimal("100.00")

    markdown = (report_dir / "current_state_report.md").read_text(encoding="utf-8")
    assert "# Отчет по текущим взаиморасчетам" in markdown
    assert "Какие покупки включены: Открытые покупки" in markdown
    assert "Доля оплат, %" in markdown
    assert "Объем расходов на человека" in markdown
    assert "Итог: + получит, − должен" in markdown
    assert "Итоговый баланс" not in markdown
    assert "Положительный итоговый баланс означает" in markdown
    assert markdown.index("## Итоговые переводы") < markdown.index("## Расходы по категориям")
    html = (report_dir / "current_state_dashboard.html").read_text(encoding="utf-8")
    assert '<meta charset="UTF-8">' in html
    assert "Отчет по текущим взаиморасчетам" in html
    assert "Доля оплат, %" in html
    assert "Объем расходов на человека" in html
    assert "Итог: + получит, − должен" in html
    assert "Итоговый баланс" not in html
    assert "Положительный итоговый баланс означает" in html
    assert "payer_total" not in html
    assert "payer_rank" not in html
    assert html.index("Итоговые переводы") < html.index("Покупки")

    workbook = load_workbook(report_dir / "current_state.xlsx")
    assert workbook.sheetnames == [
        "Сводка",
        "Итоговые переводы",
        "Покупки",
        "По категориям",
        "По плательщикам",
        "Объем расходов на человека",
        "Балансы участников",
        "Предупреждения",
        "Графики",
    ]
    purchases_sheet = workbook["Покупки"]
    purchase_headers = [cell.value for cell in purchases_sheet[3]]
    assert "ID" not in purchase_headers
    assert "payer_total" not in purchase_headers
    assert "payer_rank" not in purchase_headers
    charts_sheet = workbook["Графики"]
    assert any(
        cell.value == "Объем расходов на человека"
        for row in charts_sheet.iter_rows()
        for cell in row
    )
    payer_sheet = workbook["По плательщикам"]
    assert "Доля оплат, %" in [cell.value for cell in payer_sheet[3]]
    participant_sheet = workbook["Объем расходов на человека"]
    assert "Доля расходов" not in [cell.value for cell in participant_sheet[3]]
    assert "Объем расходов на человека" in [cell.value for cell in participant_sheet[3]]
    balances_sheet = workbook["Балансы участников"]
    balance_headers = [cell.value for cell in balances_sheet[3]]
    assert "Доля" not in balance_headers
    assert "Баланс" not in balance_headers
    assert "Объем расходов на человека" in balance_headers
    assert "Итоговый баланс" not in balance_headers
    assert "Итог: + получит, − должен" in balance_headers


def test_empty_open_scope_writes_metadata_summary_and_warning(tmp_path):
    dataset = build_current_report_dataset(sample_participants(), [], scope="open")

    report_dir = generate_current_report(dataset, tmp_path, "all")

    warnings_text = (report_dir / "tables" / "warnings.csv").read_text(encoding="utf-8-sig")
    assert EMPTY_OPEN_SCOPE_MESSAGE in warnings_text
    assert list((report_dir / "charts").glob("*.png")) == []
    metadata = json.loads((report_dir / "metadata.json").read_text(encoding="utf-8"))
    assert metadata["summary"]["purchase_count"] == 0
    assert metadata["warning_count"] >= 1


def test_cli_current_report_generates_requested_outputs(tmp_path):
    data_dir = tmp_path / "data"
    output_root = tmp_path / "reports"
    initialize_data_files(data_dir, [], [])
    participants = sample_participants()
    save_participants(data_dir / "participants.yaml", participants)
    save_purchases(data_dir / "purchases.yaml", sample_purchases())

    result = runner.invoke(
        app,
        [
            "current-report",
            "--scope",
            "open",
            "--format",
            "html",
            "--output-root",
            str(output_root),
            "--data-dir",
            str(data_dir),
        ],
    )

    assert result.exit_code == 0
    assert "Current state report generated" in result.stdout
    report_dirs = list(output_root.glob("open_*"))
    assert len(report_dirs) == 1
    assert (report_dirs[0] / "current_state_dashboard.html").is_file()
    assert not (report_dirs[0] / "current_state.xlsx").exists()
