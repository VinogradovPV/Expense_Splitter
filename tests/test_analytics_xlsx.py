from datetime import date
from decimal import Decimal
from zipfile import ZipFile

from openpyxl import load_workbook
from typer.testing import CliRunner

from expense_splitter.analytics import build_analytics_dataset, parse_period
from expense_splitter.analytics_reporting import generate_analytics_report
from expense_splitter.cli import app
from expense_splitter.defaults import DEFAULT_GROUPS, DEFAULT_PARTICIPANTS
from expense_splitter.models import Participant, Purchase
from expense_splitter.storage import initialize_data_files, save_purchases

runner = CliRunner()
EXPECTED_SHEETS = [
    "Summary",
    "Purchases",
    "By Category",
    "By Payer",
    "By Participant",
    "Balances",
    "Settlements",
    "Top Purchases",
    "Warnings",
    "Charts",
]


def dataset(with_purchase: bool = True):
    purchases = []
    if with_purchase:
        purchases.append(
            Purchase(
                id="p1",
                date=date(2026, 6, 10),
                amount=Decimal("120.50"),
                payer="Алиса",
                participants=["Алиса", "Боб"],
                purchase_name="Обед",
                category="Еда",
            )
        )
    return build_analytics_dataset(
        [Participant("Алиса"), Participant("Боб")],
        [],
        purchases,
        parse_period("month", 2026, month=6),
    )


def test_xlsx_report_contains_required_sheets_and_russian_data(tmp_path):
    report_dir = generate_analytics_report(dataset(), tmp_path, "xlsx")
    path = report_dir / "expense_analytics_2026-06.xlsx"
    workbook = load_workbook(path)

    assert workbook.sheetnames == EXPECTED_SHEETS
    summary = workbook["Summary"]
    assert summary["A1"].value == "Сводка аналитики"
    assert any(cell.value == "ID периода" for row in summary.iter_rows() for cell in row)
    assert any(cell.value == "2026-06" for row in summary.iter_rows() for cell in row)
    assert any(cell.value == 120.5 for row in summary.iter_rows() for cell in row)

    purchases = workbook["Purchases"]
    headers = [cell.value for cell in purchases[3]]
    values = [cell.value for cell in purchases[4]]
    assert "Наименование покупки" in headers
    assert values[headers.index("Наименование покупки")] == "Обед"
    assert values[headers.index("Категория")] == "Еда"
    assert values[headers.index("Плательщик")] == "Алиса"
    assert values[headers.index("Участники")] == "Алиса, Боб"
    amount_cell = purchases.cell(4, headers.index("Сумма") + 1)
    assert amount_cell.value == 120.5
    assert amount_cell.number_format == "#,##0.00"
    assert purchases.freeze_panes == "A4"
    assert purchases.auto_filter.ref
    assert len(workbook["Charts"]._images) == 6
    with ZipFile(path) as archive:
        assert not any(name.startswith("xl/tables/") for name in archive.namelist())


def test_xlsx_empty_period_has_warnings_and_chart_message(tmp_path):
    report_dir = generate_analytics_report(dataset(False), tmp_path, "xlsx")
    workbook = load_workbook(report_dir / "expense_analytics_2026-06.xlsx")

    assert "Warnings" in workbook.sheetnames
    assert workbook["Warnings"].max_row >= 4
    assert workbook["Charts"]["A3"].value == ("Графики не созданы: нет данных за выбранный период.")


def test_all_format_includes_xlsx(tmp_path):
    report_dir = generate_analytics_report(dataset(), tmp_path, "all")
    assert (report_dir / "expense_analytics_2026-06.xlsx").is_file()


def test_cli_xlsx_format(tmp_path):
    data_dir = tmp_path / "data"
    output_root = tmp_path / "reports"
    initialize_data_files(data_dir, DEFAULT_PARTICIPANTS, DEFAULT_GROUPS)
    save_purchases(
        data_dir / "purchases.yaml",
        [
            Purchase(
                id="p1",
                date=date(2026, 6, 10),
                amount=Decimal("120.50"),
                payer=DEFAULT_PARTICIPANTS[0].name,
                participants=[item.name for item in DEFAULT_PARTICIPANTS[:2]],
                purchase_name="Обед",
            )
        ],
    )

    result = runner.invoke(
        app,
        [
            "analytics",
            "--period",
            "month",
            "--year",
            "2026",
            "--month",
            "6",
            "--format",
            "xlsx",
            "--output-root",
            str(output_root),
            "--data-dir",
            str(data_dir),
        ],
    )

    assert result.exit_code == 0, result.stdout
    path = output_root / "2026" / "2026-06" / "expense_analytics_2026-06.xlsx"
    assert path.is_file()
    assert load_workbook(path).sheetnames == EXPECTED_SHEETS
