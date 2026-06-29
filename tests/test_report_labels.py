from openpyxl import load_workbook

from expense_splitter.current_report import build_current_report_dataset, generate_current_report
from expense_splitter.settlement_period_report import (
    build_settlement_period_report_dataset,
    generate_settlement_period_report,
)
from tests.test_current_report import sample_participants, sample_purchases
from tests.test_settlement_period_report import sample_period


def test_current_report_visual_summary_uses_russian_labels(tmp_path):
    dataset = build_current_report_dataset(sample_participants(), sample_purchases(), scope="open")
    report_dir = generate_current_report(dataset, tmp_path, "all")

    markdown = (report_dir / "current_state_report.md").read_text(encoding="utf-8")
    html = (report_dir / "current_state_dashboard.html").read_text(encoding="utf-8")
    summary_text = markdown + "\n" + html

    assert "Какие покупки включены" in summary_text
    assert "Открытые покупки" in summary_text
    assert "Дата формирования отчета" in summary_text
    assert "Scope:" not in summary_text
    assert "Snapshot:" not in summary_text
    assert "Scope: open" not in summary_text

    workbook = load_workbook(report_dir / "current_state.xlsx")
    assert "Сводка" in workbook.sheetnames
    assert "Графики" in workbook.sheetnames


def test_settlement_period_report_visual_summary_uses_russian_labels(tmp_path):
    dataset = build_settlement_period_report_dataset(
        [sample_period()],
        sample_purchases(),
        sample_participants(),
        "settlement_2026_06_30_001",
    )
    report_dir = generate_settlement_period_report(dataset, tmp_path, "all")

    markdown = (report_dir / "settlement_period_report.md").read_text(encoding="utf-8")
    html = (report_dir / "settlement_period_dashboard.html").read_text(encoding="utf-8")
    summary_text = markdown + "\n" + html

    assert "ID периода взаиморасчетов" in summary_text
    assert "Статус: Закрыт" in summary_text
    assert "Дата формирования отчета" in summary_text
    assert "Settlement period ID:" not in summary_text
    assert "Snapshot:" not in summary_text

    workbook = load_workbook(report_dir / "settlement_period.xlsx")
    assert "Сводка" in workbook.sheetnames
    assert "Итоговые переводы" in workbook.sheetnames
